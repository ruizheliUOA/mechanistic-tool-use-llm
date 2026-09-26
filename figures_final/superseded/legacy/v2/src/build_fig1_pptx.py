"""Build Figure 1 as a native, fully editable PowerPoint deck.

Every element is a real PPTX shape: rounded rectangles, straight connectors with
XML-injected arrowheads, ovals, and text frames. Nothing is rasterised. The slide
is sized to the ICLR text width (5.5 in) so the figure is designed at final
insertion scale and needs no downscaling.

Topology drawn here is the verified one (see FIG1_TOPOLOGY_VERIFICATION.md):
L_inj = 21 executes before L_obs = 26, so diagnosis and intervention are two
separate forward passes with the router scoring cached activations in between.
"""
from __future__ import annotations
import os, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from lxml import etree

# ----------------------------------------------------------------- palette
INK   = RGBColor(0x1B, 0x1F, 0x26)
SLATE = RGBColor(0x33, 0x47, 0x5B)
BLUE  = RGBColor(0x3D, 0x6E, 0x9C)
TEAL  = RGBColor(0x2F, 0x7D, 0x6A)
VERM  = RGBColor(0xA8, 0x45, 0x3B)
GREY  = RGBColor(0x8B, 0x91, 0x99)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SLATE_F = RGBColor(0xED, 0xF1, 0xF5)
BLUE_F  = RGBColor(0xE7, 0xEE, 0xF5)
TEAL_F  = RGBColor(0xE6, 0xF0, 0xED)
VERM_F  = RGBColor(0xF6, 0xE9, 0xE7)
GREY_F  = RGBColor(0xF0, 0xF2, 0xF4)
HAIR    = RGBColor(0xC3, 0xCA, 0xD1)

FONT = "Arial"
W, H = 5.5, 3.34

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
slide = prs.slides.add_slide(prs.slide_layouts[6])
SH = slide.shapes

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _ln(shape):
    spPr = shape._element.spPr
    ln = spPr.find(A + "ln")
    if ln is None:
        ln = etree.SubElement(spPr, A + "ln")
    return ln


def arrowhead(shape, head=False, tail=True, size="sm"):
    ln = _ln(shape)
    for tag, on in ((A + "headEnd", head), (A + "tailEnd", tail)):
        for old in ln.findall(tag):
            ln.remove(old)
        if on:
            e = etree.SubElement(ln, tag)
            e.set("type", "triangle"); e.set("w", size); e.set("len", size)


def style_text(tf, size, color, bold=False, align=PP_ALIGN.CENTER,
               anchor=MSO_ANCHOR.MIDDLE, italic=False):
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for p in tf.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.font.size, r.font.name = Pt(size), FONT
            r.font.color.rgb, r.font.bold, r.font.italic = color, bold, italic


def rect(x, y, w, h, text="", *, fill=None, line=SLATE, lw=0.6, size=6,
         tc=None, bold=False, radius=0.12, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = SH.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    try:
        s.adjustments[0] = radius
    except Exception:
        pass
    s.shadow.inherit = False
    s.text_frame.text = text
    style_text(s.text_frame, size, tc or INK, bold)
    return s


def tb(x, y, w, h, text, *, size=6, color=INK, bold=False,
       align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, italic=False):
    t = SH.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.text_frame.text = text
    style_text(t.text_frame, size, color, bold, align, anchor, italic)
    return t


def conn(x1, y1, x2, y2, *, color=INK, lw=0.6, arrow=True, dash=None):
    c = SH.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                         Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw)
    if dash:
        ln = _ln(c)
        for old in ln.findall(A + "prstDash"):
            ln.remove(old)
        etree.SubElement(ln, A + "prstDash").set("val", dash)
    arrowhead(c, tail=arrow)
    return c


def elbow(pts, *, color=INK, lw=0.6):
    """Orthogonal polyline as separate editable connectors; arrow on the last."""
    out = []
    for i in range(len(pts) - 1):
        (x1, y1), (x2, y2) = pts[i], pts[i + 1]
        out.append(conn(x1, y1, x2, y2, color=color, lw=lw,
                        arrow=(i == len(pts) - 2)))
    return out


def oval(cx, cy, r, *, line=BLUE, fill=WHITE, lw=0.6):
    s = SH.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r),
                     Inches(2 * r), Inches(2 * r))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def stack(x, y, w, h, n, mark):
    gap = 0.016
    bw = (w - gap * (n - 1)) / n
    cs = []
    for i in range(n):
        xi = x + i * (bw + gap)
        hit = i == mark
        rect(xi, y, bw, h, fill=BLUE_F if hit else SLATE_F,
             line=BLUE if hit else SLATE, lw=0.9 if hit else 0.5, radius=0.18)
        cs.append(xi + bw / 2)
    return cs


def dashed_panel(x, y, w, h, title="", tc=SLATE):
    s = rect(x, y, w, h, fill=None, line=GREY, lw=0.5, radius=0.03,
             shape=MSO_SHAPE.RECTANGLE)
    ln = _ln(s)
    etree.SubElement(ln, A + "prstDash").set("val", "dash")
    if title:
        tb(x + 0.04, y - 0.045, w - 0.08, 0.09, title, size=6, color=tc,
           bold=True, align=PP_ALIGN.LEFT)
    return s


# ================================================================== OFFLINE
dashed_panel(0.05, 0.10, 3.24, 0.46, "OFFLINE  ·  channel discovery")
rect(0.11, 0.20, 0.55, 0.28, "baseline\npreds + gold", fill=GREY_F, line=GREY, size=5.5)
conn(0.66, 0.34, 0.74, 0.34, color=GREY)
rect(0.75, 0.20, 0.58, 0.28, "directed error\ntopology", fill=GREY_F, line=GREY, size=5.5)
conn(1.33, 0.34, 1.41, 0.34, color=GREY)
rect(1.42, 0.20, 0.62, 0.28, "channel\nc = (g → s)", fill=BLUE_F, line=BLUE, size=5.5)
conn(2.04, 0.28, 2.14, 0.25, color=BLUE)
conn(2.04, 0.40, 2.14, 0.43, color=BLUE)
rect(2.15, 0.15, 0.52, 0.18, "router  r_c", fill=BLUE_F, line=BLUE, size=5.5)
rect(2.15, 0.36, 0.52, 0.18, "direction  d_c", fill=BLUE_F, line=BLUE, size=5.5)
tb(2.70, 0.20, 0.56, 0.28, "discovered from\nthe model's own\nerrors", size=5,
   color=GREY, italic=True)

# ================================================================== PASS 1
tb(0.05, 0.60, 1.5, 0.10, "① DIAGNOSE  ·  read", size=6.5, color=SLATE,
   bold=True, align=PP_ALIGN.LEFT)
rect(0.05, 0.79, 0.46, 0.30, "request\n+ tools", size=5.5)
conn(0.51, 0.94, 0.57, 0.94)
c1 = stack(0.58, 0.79, 1.42, 0.30, 9, 6)
tb(0.58, 1.09, 0.30, 0.09, "B₁", size=5, color=SLATE, align=PP_ALIGN.LEFT)
tb(c1[6] - 0.16, 1.09, 0.32, 0.09, "L_obs", size=5, color=BLUE, bold=True)
tb(0.76, 1.09, 0.50, 0.09, "θ frozen", size=5, color=SLATE, align=PP_ALIGN.LEFT)
elbow([(c1[6], 0.79), (c1[6], 0.72), (2.34, 0.72), (2.34, 0.79)], color=BLUE)
tb(c1[6] + 0.04, 0.63, 0.40, 0.09, "h_obs", size=5.5, color=BLUE,
   align=PP_ALIGN.LEFT)
rect(2.10, 0.79, 0.48, 0.30, "router\nr_c(h_obs)", fill=BLUE_F, line=BLUE, size=5.5)
conn(2.58, 0.94, 2.64, 0.94, color=BLUE)
rect(2.65, 0.79, 0.40, 0.30, "p_c ≥ τ ?", line=BLUE, size=5.5)
tb(2.10, 1.12, 0.95, 0.09, "cached, scored offline", size=5,
   color=GREY, italic=True)

# ================================================================== PASS 2
tb(0.05, 1.24, 1.7, 0.10, "② INTERVENE  ·  gated write", size=6.5, color=SLATE,
   bold=True, align=PP_ALIGN.LEFT)
conn(0.28, 1.09, 0.28, 1.55, color=GREY, dash="dash")
tb(0.05, 1.56, 0.46, 0.09, "same input", size=5, color=GREY)
conn(0.51, 1.82, 0.57, 1.82)
c2 = stack(0.58, 1.67, 1.42, 0.30, 9, 3)
tb(c2[3] - 0.16, 1.97, 0.32, 0.09, "L_inj", size=5, color=BLUE, bold=True)
tb(1.70, 1.97, 0.30, 0.09, "B_L", size=5, color=SLATE, align=PP_ALIGN.RIGHT)
oval(c2[3], 1.56, 0.055)
conn(c2[3] - 0.035, 1.56, c2[3] + 0.035, 1.56, color=BLUE, arrow=False)
conn(c2[3], 1.525, c2[3], 1.595, color=BLUE, arrow=False)
tb(c2[3] + 0.08, 1.51, 0.72, 0.10, "h′ = h + α d_c", size=6, color=BLUE,
   align=PP_ALIGN.LEFT)
elbow([(2.85, 1.09), (2.85, 1.38), (c2[3], 1.38), (c2[3], 1.50)], color=BLUE)
tb(1.62, 1.31, 0.62, 0.09, "fire", size=5, color=BLUE, align=PP_ALIGN.RIGHT)
conn(2.00, 1.82, 2.08, 1.82)
rect(2.09, 1.67, 0.42, 0.30, "â′", line=INK, size=7)
tb(0.05, 2.05, 2.46, 0.09,
   "L_inj executes before L_obs — two passes over the same frozen model",
   size=5, color=GREY, italic=True, align=PP_ALIGN.LEFT)

# ================================================================== DESTINATIONS
dashed_panel(3.36, 0.10, 2.09, 1.16, "DESTINATION-RESOLVED OUTCOME")
elbow([(2.51, 1.82), (3.00, 1.82), (3.00, 0.68), (3.44, 0.68)])
oval(3.46, 0.68, 0.017, line=INK, fill=INK)
rows = [(0.24, GREY, GREY_F, "SOURCE_RETAINED", "â′ = s", "still wrong"),
        (0.58, TEAL, TEAL_F, "GOLD_ARRIVAL", "â′ = g", "correct"),
        (0.92, VERM, VERM_F, "OTHER_WRONG", "â′ = w ∉ {g, s}", "still wrong")]
for yy, col, fc, nm, ms, vd in rows:
    rect(3.62, yy, 1.16, 0.26, "", fill=fc, line=col, lw=0.6, radius=0.14)
    tb(3.66, yy + 0.015, 1.10, 0.11, nm, size=5.5, color=col, bold=True,
       align=PP_ALIGN.LEFT)
    tb(3.66, yy + 0.135, 1.10, 0.11, ms, size=5.5, align=PP_ALIGN.LEFT)
    conn(3.48, 0.68, 3.61, yy + 0.13, color=col, lw=0.5)
    conn(4.78, yy + 0.13, 4.86, yy + 0.13, color=col, lw=0.5, arrow=False)
    tb(4.88, yy + 0.07, 0.54, 0.11, vd, size=5, color=col,
       bold=(col == TEAL), align=PP_ALIGN.LEFT)

# ================================================================== PRESERVATION
dashed_panel(3.36, 1.38, 2.09, 0.68, "SECOND POPULATION  ·  preservation")
elbow([(3.00, 1.82), (3.00, 1.72), (3.42, 1.72)])
rect(3.44, 1.58, 0.72, 0.28, "baseline-correct,\nexposed", size=5.5)
conn(4.16, 1.66, 4.30, 1.58, color=TEAL, lw=0.5)
conn(4.16, 1.78, 4.30, 1.86, color=VERM, lw=0.5)
rect(4.31, 1.48, 0.62, 0.20, "retained", fill=TEAL_F, line=TEAL, tc=TEAL, size=5.5)
rect(4.31, 1.78, 0.62, 0.20, "broken", fill=VERM_F, line=VERM, tc=VERM, size=5.5)
tb(4.96, 1.58, 0.46, 0.28, "a distinct\nestimand", size=5, color=GREY, italic=True)

# ================================================================== CLAIM LAYER
conn(0.05, 2.17, 5.45, 2.17, color=HAIR, lw=0.5, arrow=False)
tb(0.05, 2.22, 3.0, 0.11, "what was shown, and what may be asserted",
   size=6.5, color=SLATE, bold=True, align=PP_ALIGN.LEFT)

rect(0.05, 2.44, 0.74, 0.28, "Adjudicable", line=SLATE, lw=1.0, size=6.5)
tb(0.05, 2.74, 0.74, 0.10, "entry condition", size=5, color=GREY)
conn(0.79, 2.58, 0.90, 2.58, color=SLATE)

rect(0.91, 2.38, 2.34, 0.34, "", fill=BLUE_F, line=BLUE, lw=0.6, radius=0.06)
tb(0.91, 2.28, 2.34, 0.09, "SCIENTIFIC PROPERTY LAYER", size=5, color=BLUE, bold=True)
for i, (nm, sub) in enumerate((("Readable", "error is decodable"),
                               ("Steerable", "effect is direction-specific"),
                               ("Correctable", "movement lands on g"))):
    xx = 0.96 + i * 0.77
    rect(xx, 2.44, 0.72, 0.28, nm, fill=WHITE, line=BLUE, size=6.5)
    if i:
        conn(xx - 0.05, 2.58, xx - 0.01, 2.58, color=BLUE)
tb(0.91, 2.74, 2.34, 0.09, "a property claim about the tested setting",
   size=5, color=BLUE, italic=True)
conn(3.25, 2.58, 3.36, 2.58, color=SLATE)

rect(3.37, 2.40, 1.30, 0.36, "", line=SLATE, lw=1.4, radius=0.10)
rect(3.40, 2.43, 1.24, 0.30, "Licensable", line=SLATE, lw=0.5, size=6.5)
tb(3.37, 2.28, 1.30, 0.09, "EVIDENTIAL DECISION", size=5, color=SLATE, bold=True)
tb(4.70, 2.42, 0.74, 0.32, "ADMIT\nDECLINE\nNO-GO", size=5.5, align=PP_ALIGN.LEFT)
tb(3.37, 2.75, 1.30, 0.09, "whether the evidence justifies that claim",
   size=5, color=SLATE, italic=True)
tb(3.37, 2.85, 2.07, 0.09, "DECLINE ≠ intrinsically uncorrectable",
   size=5, color=VERM, italic=True)

crit = ["Specificity", "Destination", "Preservation", "Yield", "Evidence"]
for i, nm in enumerate(crit):
    xx = 0.72 + i * 0.72
    rect(xx, 3.02, 0.66, 0.20, nm, fill=GREY_F, line=GREY, size=5.5)
    conn(xx + 0.33, 3.02, xx + 0.33, 2.96, color=GREY, lw=0.5, arrow=False)
conn(1.05, 2.96, 4.29, 2.96, color=GREY, lw=0.5, arrow=False)
elbow([(4.02, 2.96), (4.02, 2.78)], color=GREY, lw=0.5)
tb(0.05, 3.04, 0.62, 0.16, "EVIDENTIAL\nQUALIFICATION", size=5, color=GREY,
   bold=True, align=PP_ALIGN.LEFT)

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "final")
os.makedirs(out, exist_ok=True)
path = os.path.join(out, "fig1_sakiko_core_FINAL.pptx")
prs.save(path)
print("shapes:", len(SH._spTree.findall(
    "{http://schemas.openxmlformats.org/presentationml/2006/main}sp")) )
print("saved", path)
